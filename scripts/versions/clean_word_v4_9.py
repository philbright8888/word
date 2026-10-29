#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""WPS/PDF -> DOCX cleaner V4.9 grid-consistent bank columns + Notes reflow + fixed bank widths. Python 3.12+, pip install python-docx.

Put this script beside .docx files, run it and select a file number.
No translations, no blanket deletion of English data; originals are never overwritten.
"""
from __future__ import annotations
from pathlib import Path
from collections import Counter
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
FONT_PT = 8
LINE_PT = 8                   # fixed 8-point line spacing (may clip some glyphs)
# Exact widths in cm supplied by user. Notes is the remainder of 27.81 cm.
TOTAL_TABLE_CM = 27.81
FIXED_CM = [0.60, 1.00, 0.97, 1.25, 0.75, 4.75, 2.00, 2.00, 2.25, 3.25]
BANK_COL_CM = FIXED_CM + [round(TOTAL_TABLE_CM - sum(FIXED_CM), 6)]
assert len(BANK_COL_CM) == 11 and abs(sum(BANK_COL_CM)-TOTAL_TABLE_CM)<1e-7
NEW_PAGE_FOR_EACH_IMAGE = True  # second and subsequent BODY inline-image paragraphs
HEADERS = ["No.", "Bk.D.", "Val.D.", "Type", "Vou.",
           "Vou.No./Trans.No./Details", "DebitAmount", "CreditAmount",
           "Balance", "ReferenceNo.", "Notes"]
DATA_ALIGNS = [WD_ALIGN_PARAGRAPH.CENTER]*5 + [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.RIGHT]*3 + [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]
STATS = Counter()
WARNINGS = []


def get_or_add(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el


def set_val(parent, tag, value):
    get_or_add(parent, tag).set(qn('w:val'), str(value))


def contains_image(p):
    return bool(p._p.xpath('.//w:drawing | .//w:pict | .//w:object'))


def preserve_p(p):
    return contains_image(p) or bool(p._p.xpath('.//w:sectPr | .//w:fldChar | .//w:bookmarkStart'))


def set_run_font(run):
    rpr = run._element.get_or_add_rPr()
    rf = get_or_add(rpr, 'w:rFonts')
    for attr, value in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),
                        ('eastAsia','SimSun'),('cs','Times New Roman')]:
        rf.set(qn('w:'+attr), value)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','csTheme'):
        rf.attrib.pop(qn('w:'+key), None)
    run.font.size = Pt(FONT_PT)
    set_val(rpr,'w:sz',FONT_PT*2)
    set_val(rpr,'w:szCs',FONT_PT*2)
    set_val(rpr,'w:w',100)                # character width 100%
    set_val(rpr,'w:spacing',0)            # 0 twips tracking
    set_val(rpr,'w:position',0)
    for tag in ('w:fitText','w:kern','w:vertAlign','w:smallCaps','w:caps'):
        node = rpr.find(qn(tag))
        if node is not None: rpr.remove(node)
    STATS['runs'] += 1


def normalize_spaces(s):
    # Leave text content intact except extra whitespace and formatting artifacts.
    return re.sub(r'[ \t\u00a0\u3000\u2002-\u2009]+',' ',s)


def clean_runs(p):
    runs = p.runs
    for run in runs: set_run_font(run)
    # Avoid touching mixed drawing/field/hyperlink XML by rewriting run strings.
    if p._p.xpath('.//w:drawing | .//w:pict | .//w:object | .//w:fldChar | .//w:instrText | .//w:hyperlink'):
        return
    if not runs: return
    if any('\n' in r.text or '\r' in r.text for r in runs):
        # Actual line breaks are preserved in regular text (Notes handled separately).
        return
    original = ''.join(r.text for r in runs)
    desired = normalize_spaces(original).strip(' ')
    if original == desired: return
    # Keep each nonspace sequence in its original styled run.
    has_text = False
    for run in runs:
        t = normalize_spaces(run.text)
        if not has_text: t = t.lstrip(' ')
        if t.strip(): has_text = True
        run.text = t
    for run in reversed(runs):
        if run.text:
            run.text = run.text.rstrip(' ')
            break
    prior = None
    for run in runs:
        if prior is not None and prior.text.endswith(' ') and run.text.startswith(' '):
            run.text = run.text.lstrip(' ')
        if run.text: prior = run
    STATS['whitespace_cleaned'] += 1


def clean_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, keep_image_layout=True):
    pic = contains_image(p)
    pf = p.paragraph_format
    if not (pic and keep_image_layout):
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(LINE_PT)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if pic and keep_image_layout:
        # The user's inline images are at the top right. Keep anchoring and indentation.
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        STATS['protected_image_paragraphs'] += 1
    else:
        pf.left_indent = Pt(0)
        pf.right_indent = Pt(0)
        pf.first_line_indent = Pt(0)
        p.alignment = alignment
    pp = p._p.get_or_add_pPr()
    if not (pic and keep_image_layout):
        sp = get_or_add(pp,'w:spacing')
        for name, v in [('before',0),('after',0),('line',LINE_PT*20),
                        ('beforeAutospacing',0),('afterAutospacing',0)]:
            sp.set(qn('w:'+name),str(v))
        sp.set(qn('w:lineRule'),'exact')
        set_val(pp,'w:contextualSpacing',0)
    clean_runs(p)
    STATS['paragraphs'] += 1


def safe_empty(p):
    if p.text.strip(): return False
    if preserve_p(p): return False
    if p._p.xpath('.//w:br | .//w:cr | .//w:tab | .//w:footnoteReference | .//w:endnoteReference'):
        return False
    return True


def remove_blank_paragraphs(paragraphs, in_cell=False, protect_header_area=False):
    # Optional protection of page-head paragraphs to avoid shifting image/account/table.
    for p in list(paragraphs):
        if not safe_empty(p): continue
        prev = p._p.getprevious()
        next_el = p._p.getnext()
        if protect_header_area and (prev is not None and prev.xpath('.//w:drawing') or
                                    next_el is not None and next_el.xpath('.//w:drawing')):
            continue
        parent = p._p.getparent()
        if parent is None: continue
        if in_cell:
            blocks = [n for n in parent if n.tag in (qn('w:p'),qn('w:tbl'))]
            if len(blocks)<=1 or (blocks[-1] is p._p and any(n.tag==qn('w:tbl') for n in blocks)):
                continue
        parent.remove(p._p)
        STATS['empty_paragraphs_removed'] += 1


def compact_cell(cell):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr = cell._tc.get_or_add_tcPr()
    mar = get_or_add(tcpr, 'w:tcMar')
    for side in ('top','bottom','left','right','start','end'):
        n = get_or_add(mar,'w:'+side)
        n.set(qn('w:w'),'0')
        n.set(qn('w:type'),'dxa')
    STATS['cells'] += 1


def cell_text(cell):
    return ' '.join(p.text for p in cell.paragraphs).strip()


def normalized_header(s):
    return re.sub(r'[^a-z0-9]+','',s.casefold())


# Exact source headings supplied by the user, plus their standardized forms.
# Match actual text rather than assuming the PDF-converted table has 11 grid columns.
HEADER_FORMS = [
    ['Serial No. No.', 'Serial No.', 'No.'],
    ['Booked Date Bk.D.', 'Booked Date', 'Bk.D.'],
    ['Value Date Vou.', 'Value   Date Vou.', 'Value Date Val.D.', 'Value Date', 'Val.D.'],
    ['Transaction Type Type', 'Transaction Type', 'Type'],
    ['Debit Vou.', 'Vou.'],
    ['Vou. No./Trans. No./Details   Vou. No./Trans. No./Details',
     'Vou. No./Trans. No./Details', 'Vou.No./Trans.No./Details'],
    ['Debit Amount Debit Amount', 'Debit Amount', 'DebitAmount'],
    ['Credit Amount Credit Amount', 'Credit Amount', 'CreditAmount'],
    ['Balance Balance', 'Balance'],
    ['Institution/Cashier/Sequence Reference No.',
     'Institution/Cashier/Sequence   Reference No.', 'Reference No.', 'ReferenceNo.'],
    ['Notes Notes', 'Notes'],
]
HEADER_MAP = {}
for _j, _variants in enumerate(HEADER_FORMS):
    for _v in _variants:
        HEADER_MAP[normalized_header(_v)] = _j


def header_index(text):
    # Full-cell matching only: never rewrite ordinary transaction details.
    return HEADER_MAP.get(normalized_header(text))


def unique_row_cells(row):
    # python-docx's row.cells repeats the SAME merged cell at several grid positions.
    seen = set()
    for grid_col, cell in enumerate(row.cells):
        key = id(cell._tc)
        if key not in seen:
            seen.add(key)
            yield grid_col, cell


def detect_header(table, label):
    # Find recognizable heading text in the first few physical rows.
    # A WPS table may use 16 columns with merged cells, not 11.
    hits = {}
    for ri, row in enumerate(table.rows[:min(len(table.rows), 6)]):
        for grid_col, c in unique_row_cells(row):
            j = header_index(cell_text(c))
            if j is not None:
                hits.setdefault(j, []).append((ri, grid_col, c))
    # Require several distinctive anchors, not an incidental 'Notes' or 'Balance'.
    anchors = {0, 1, 5, 6, 7, 9, 10}
    if len(hits) < 5 or len(set(hits).intersection(anchors)) < 3:
        return None, {}
    for j, places in hits.items():
        if len(places) > 1:
            WARNINGS.append(f'{label}: heading {HEADERS[j]} appears {len(places)} times; only unique cell matches are changed.')
    STATS['bank_tables_detected_by_text'] += 1
    return max(ri for places in hits.values() for ri, _, _ in places), hits


def set_cell_text_preserving_cell(cell, s):
    # Only for recognized header cells; retains cell properties/borders.
    para = cell.paragraphs[0]
    for run in para.runs: run.text=''
    if para.runs: para.runs[0].text=s
    else: para.add_run(s)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    clean_paragraph(para,WD_ALIGN_PARAGRAPH.CENTER)


def regroup_notes(cell):
    """Reflow *within one Notes cell*, never across table rows.

    WPS PDF conversion may insert paragraph breaks, explicit line breaks, tabs,
    and multiple runs. Consolidate the text into a single paragraph and let
    Word perform wrapping at the actual Notes column width (8.99 cm).
    """
    if cell.tables:
        STATS['notes_skipped_nested_table'] += 1
        return
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        return
    # Nontext objects must never be destroyed by a text-only reflow.
    unsafe = './/w:drawing | .//w:pict | .//w:object | .//w:fldChar | .//w:instrText | .//w:hyperlink | .//w:bookmarkStart | .//w:bookmarkEnd | .//w:footnoteReference | .//w:endnoteReference'
    if any(p._p.xpath(unsafe) for p in paragraphs):
        STATS['notes_skipped_complex_content'] += 1
        return

    # Both paragraph boundaries and manual line breaks from PDF are treated as
    # unwanted wraps. Preserve a word separator; don't join two transactions.
    pieces = []
    for p in paragraphs:
        # Paragraph.text exposes explicit w:br/w:cr as newline, and w:tab as tab.
        part = re.sub(r'[\r\n\t]+', ' ', p.text)
        part = normalize_spaces(part).strip()
        if part:
            pieces.append(part)
    if not pieces:
        return
    joined = normalize_spaces(' '.join(pieces)).strip()
    if not joined:
        return

    previous = '\n'.join(p.text for p in paragraphs)
    changed = len(paragraphs) != 1 or previous != joined
    if changed:
        first = paragraphs[0]
        # Retain the first paragraph object (and cell/border structure).
        # Remove textual runs, then insert the joined plain text as one run.
        # The document's font and size normalization applies to the new run.
        for run in first.runs:
            run.text = ''
        if first.runs:
            first.runs[0].text = joined
        else:
            first.add_run(joined)
        for extra in paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        STATS['notes_cells_reflowed'] += 1
        STATS['notes_hard_breaks_removed'] += max(0, previous.count('\n'))

    p = cell.paragraphs[0]
    # Disable inherited special paragraph layout that may retain unintended gaps.
    pp = p._p.get_or_add_pPr()
    for tag in ('w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:tabs'):
        element = pp.find(qn(tag))
        if element is not None:
            pp.remove(element)
    clean_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT)

    # Allow ordinary Word wrapping; never squeeze whole Notes cell onto one line.
    tcpr = cell._tc.get_or_add_tcPr()
    for tag in ('w:noWrap', 'w:tcFitText'):
        element = tcpr.find(qn(tag))
        if element is not None:
            tcpr.remove(element)

def approx_width(s):
    unit=0
    for ch in s.replace('\n',' '):
        if '\u4e00'<=ch<='\u9fff':unit+=1
        elif ch in 'ilI.,:;!|\'` ':unit+=.22
        elif ch in 'mwMW@%':unit+=.8
        elif ch.isupper() or ch.isdigit():unit+=.53
        else:unit+=.45
    return max(.16,(unit+1.1)*8/72+.04)


def cell_grid_span(cell):
    pr = cell._tc.get_or_add_tcPr()
    el = pr.find(qn('w:gridSpan'))
    return max(1, int(el.get(qn('w:val'))) if el is not None else 1)


def build_partial_mapping(table, hits, label):
    """Build a complete 11-column partition of the actual Word table grid.

    Word's underlying grid may have 16+ positions even when there are 11
    visible columns. Anchor each recognized logical heading at its physical
    starting grid position; fill missing headings only if a free physical grid
    position exists between the surrounding anchors. Never leave unassigned
    grid widths, as this makes a partial table wider than 27.81 cm.
    """
    ngrid = len(table.columns)
    if ngrid < 11:
        WARNINGS.append(f'{label}: grid has only {ngrid} slots (<11); no fixed-width mapping.')
        return {}, {}

    # Each heading may be detected in several physical rows. Prefer one
    # candidate that sits in left-to-right logical order.
    candidates = {j: sorted(set(g for _, g, _ in hits.get(j, []))) for j in range(11)}
    # Candidate choice is dynamic programming: maximize number of distinct
    # matched headings with strictly increasing physical grid positions.
    states = {(-1, 0): (0, [])}
    for j in range(11):
        next_states = {}
        for (last, unused), (score, picks) in states.items():
            key = (last, unused+1)
            val = (score, picks + [None])
            if key not in next_states or next_states[key][0] < val[0]:
                next_states[key] = val
            for g in candidates[j]:
                if g <= last: continue
                key = (g, 0)
                val = (score+1, picks+[g])
                if key not in next_states or next_states[key][0] < val[0]:
                    next_states[key] = val
        states = next_states
    score, picks = max(states.values(), key=lambda v: v[0])
    if score < 9:
        WARNINGS.append(f'{label}: only {score}/11 ordered column anchors found; width adjustment skipped.')
        return {}, {}

    # Preserve each known anchor; interpolate missing column starts only when
    # there are enough grid slots, and preserve ascending column order.
    starts = list(picks)
    for j in range(11):
        if starts[j] is not None: continue
        before = next((k for k in range(j-1,-1,-1) if starts[k] is not None), None)
        after = next((k for k in range(j+1,11) if starts[k] is not None), None)
        lower = (starts[before] + j-before) if before is not None else j
        upper = (starts[after] - (after-j)) if after is not None else (ngrid-1-(10-j))
        if lower > upper:
            WARNINGS.append(f'{label}: cannot locate missing {HEADERS[j]} between neighboring columns; width adjustment skipped.')
            return {}, {}
        starts[j] = lower
        STATS['missing_column_positions_inferred'] += 1

    if starts[0] != 0:
        # Some documents have a leading gridBefore; include it in first column.
        starts[0] = 0
    if any(starts[i] >= starts[i+1] for i in range(10)):
        WARNINGS.append(f'{label}: overlapping logical column starts; skipped.')
        return {}, {}
    segments = {j: list(range(starts[j], starts[j+1] if j<10 else ngrid)) for j in range(11)}
    mapping = {g:j for j, gs in segments.items() for g in gs}
    if sorted(mapping) != list(range(ngrid)):
        WARNINGS.append(f'{label}: incomplete underlying grid assignment; skipped.')
        return {}, {}
    return mapping, segments


def apply_fixed_bank_widths(table, hits, label, segments):
    """Allocate the full table grid as 11 logical columns totalling 27.81 cm."""
    if len(segments)!=11:
        STATS['fixed_widths_skipped'] += 1
        return False
    ngrid=len(table.columns)
    old=[]
    for gc in table._tbl.tblGrid.gridCol_lst:
        raw=gc.get(qn('w:w'))
        old.append(max(1.,float(raw)) if raw and raw.isdigit() else 1.)
    grid_cm=[0.]*ngrid
    for j, gs in segments.items():
        weights=[old[g] for g in gs]
        total=sum(weights)
        # Preserve proportions of existing subdivisions inside a logical column.
        for g,w in zip(gs,weights):
            grid_cm[g]=BANK_COL_CM[j]*w/total
        STATS['fixed_columns_adjusted'] += 1
    table.autofit=False
    tblpr=table._tbl.tblPr
    layout=get_or_add(tblpr,'w:tblLayout')
    layout.set(qn('w:type'),'fixed')
    tw=get_or_add(tblpr,'w:tblW')
    tw.set(qn('w:type'),'dxa')
    tw.set(qn('w:w'),str(round(TOTAL_TABLE_CM*1440/2.54)))
    # Grid widths and every physical cell's tcW must agree. Word otherwise
    # recalculates a different grid and can shift body content.
    for g,gc in enumerate(table._tbl.tblGrid.gridCol_lst):
        gc.set(qn('w:w'),str(round(grid_cm[g]*1440/2.54)))
    for row in table.rows:
        gi=row.grid_cols_before
        for tc in row._tr.tc_lst:
            pr=tc.get_or_add_tcPr()
            el=pr.find(qn('w:gridSpan'))
            span=max(1,int(el.get(qn('w:val'))) if el is not None else 1)
            end=min(ngrid,gi+span)
            if gi>=ngrid: break
            tcw=get_or_add(pr,'w:tcW')
            tcw.set(qn('w:type'),'dxa')
            tcw.set(qn('w:w'),str(round(sum(grid_cm[gi:end])*1440/2.54)))
            gi=end
    STATS['bank_fixed_width_tables'] += 1
    STATS['notes_width_cm_x100']=round(BANK_COL_CM[10]*100)
    return True

def process_table(table, section, seen, label):
    last_header_row, hits = detect_header(table, label)
    header_rows = set()
    mapping, segments = {}, {}
    if hits:
        for j, places in hits.items():
            if len(places) != 1:
                continue
            ri, grid_col, c = places[0]
            set_cell_text_preserving_cell(c, HEADERS[j])
            header_rows.add(ri)
            STATS['headers_standardized'] += 1
        STATS['bank_tables'] += 1
        if len(hits) != 11:
            missing = ', '.join(HEADERS[i] for i in range(11) if i not in hits)
            WARNINGS.append(f'{label}: recognized {len(hits)}/11 heading labels; missing: {missing}. Column position may be inferred but text is preserved.')
        mapping, segments = build_partial_mapping(table, hits, label)
        if not mapping:
            STATS['bank_alignment_skipped'] += 1
    for i,row in enumerate(table.rows):
        row.height = None
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for j,c in unique_row_cells(row):
            key = id(c._tc)
            if key in seen: continue
            seen.add(key)
            compact_cell(c)
            if i <= last_header_row and hits:
                align = WD_ALIGN_PARAGRAPH.CENTER
            elif mapping and i > last_header_row:
                align = DATA_ALIGNS[mapping[j]] if j in mapping else WD_ALIGN_PARAGRAPH.LEFT
            else:
                align = WD_ALIGN_PARAGRAPH.LEFT
            if mapping and i > last_header_row and mapping.get(j) == 10:
                regroup_notes(c)
            for p in c.paragraphs: clean_paragraph(p, align)
            remove_blank_paragraphs(c.paragraphs, in_cell=True)
            for k,nested in enumerate(c.tables): process_table(nested,section,seen,label+f'/nested{k}')
    if hits:
        apply_fixed_bank_widths(table,hits,label,segments)
    else:
        STATS['unrecognized_tables_widths_preserved'] += 1
    STATS['tables'] += 1


def clear_breaks(root):
    # Deletes hard page and column breaks, including pageBreakBefore, except
    # page-start images are re-paginated in a separate pass.
    for p in root.iter(qn('w:p')):
        pr=p.find(qn('w:pPr'))
        if pr is not None:
            el=pr.find(qn('w:pageBreakBefore'))
            if el is not None:pr.remove(el); STATS['page_before_removed']+=1
            sect=pr.find(qn('w:sectPr'))
            if sect is not None:
                pr.remove(sect)
                STATS['intermediate_sections_removed']+=1
        for br in list(p.xpath('.//w:br')):
            if br.get(qn('w:type')) in ('page','column'):
                br.getparent().remove(br)
                STATS['manual_breaks_removed']+=1
        for mark in list(p.xpath('.//w:lastRenderedPageBreak')):
            mark.getparent().remove(mark)
    # Root body final w:sectPr is intentionally untouched.


def protect_page_images(doc):
    # Inline picture in BODY: place the 2nd onward at the start of a new page.
    # No repositioning/editing of the actual image drawing element is performed.
    image_paragraphs=[]
    for p in doc.paragraphs:
        if contains_image(p):
            p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            image_paragraphs.append(p)
    if NEW_PAGE_FOR_EACH_IMAGE:
        for p in image_paragraphs[1:]:
            pr=p._p.get_or_add_pPr()
            set_val(pr,'w:pageBreakBefore',1)
            STATS['image_page_breaks_added']+=1
    STATS['body_image_paragraphs']=len(image_paragraphs)
    if not image_paragraphs:
        WARNINGS.append('No body inline images detected; page-start image rule not applied (may be in headers or floating objects).')
    # A paragraph with multiple images is not split; avoids changing picture grouping.


def choose_file():
    files=sorted((p for p in BASE.glob('*.docx') if not p.name.startswith('~$')
                 and not p.stem.casefold().endswith('_cleaned')),
                 key=lambda p:p.name.casefold())
    if not files:
        print('当前目录没有待处理的 DOCX:',BASE)
        return None
    print('\n当前目录:',BASE)
    for i,f in enumerate(files,1):print(f'{i:2}. {f.name}')
    a=input('\n请输入文件编号（q 退出）：').strip()
    if a.lower()=='q':return None
    if not a.isdigit() or not 1<=int(a)<=len(files):
        print('文件编号无效');return None
    return files[int(a)-1]


def main():
    src=choose_file()
    if src is None:return
    dest=src.with_name(src.stem+'_v49_cleaned.docx')
    if dest.exists():
        print('输出文件已存在，未覆盖：',dest)
        return
    doc=Document(src)
    section=doc.sections[0]
    clear_breaks(doc.element.body)
    for p in doc.paragraphs:clean_paragraph(p)
    # Preserve blank paragraphs immediately around images in the page-head region.
    remove_blank_paragraphs(doc.paragraphs,protect_header_area=True)
    seen=set()
    for i,table in enumerate(doc.tables,1):process_table(table,section,seen,f'Table {i}')
    protect_page_images(doc)
    story_set=set()
    for s in doc.sections:
        for story in (s.header,s.footer,s.first_page_header,s.first_page_footer,
                      s.even_page_header,s.even_page_footer):
            if story.is_linked_to_previous or id(story._element) in story_set:continue
            story_set.add(id(story._element))
            # Don't remove header/footer page structure; just normalize text.
            for p in story.paragraphs:clean_paragraph(p)
            remove_blank_paragraphs(story.paragraphs,protect_header_area=True)
            for i,t in enumerate(story.tables,1):process_table(t,s,seen,f'Header/Footer table {i}')
    doc.save(dest)
    report=dest.with_suffix('.log.txt')
    report.write_text('SOURCE: '+str(src)+'\nOUTPUT: '+str(dest)+'\n\n'+
                      '\n'.join(f'{k}: {v}' for k,v in sorted(STATS.items()))+
                      '\n\nWARNINGS:\n'+'\n'.join(WARNINGS or ['None'])+
                      '\n\nNOTICE: Fixed 8pt line spacing may clip text. Precise image page coordinates '
                      'and column single-line fitting cannot be guaranteed after reflow. '
                      'Intermediate section breaks removed may change section-specific layout. '
                      'Review the resulting Word file.\n',encoding='utf-8')
    print('\n处理完成：',dest,'\n日志：',report)
    if WARNINGS:print('警告数量：',len(WARNINGS),'（见日志）')


if __name__=='__main__':
    try:main()
    except Exception as exc:
        print('处理失败：',repr(exc))
        raise