#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Word 银行对账单格式处理：保护图片、列宽和原文件。Python 3.12+，pip install python-docx。

只替换完整匹配的11个银行表头；不删除交易行，不处理表格结构，不改任何宽度。
图片段落不改段落属性，图片所在 run 也不修改。
"""
from pathlib import Path
from collections import Counter
import re
try:
    from docx import Document
except ModuleNotFoundError as exc:
    if exc.name == "docx":
        print("缺少 python-docx。请先在当前 Python 环境运行：")
        print("  python3 -m pip install python-docx")
        print("如果出现 externally-managed-environment，请创建虚拟环境：")
        print("  python3 -m venv .venv && source .venv/bin/activate && python -m pip install python-docx")
        raise SystemExit(1)
    raise
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path.cwd()  # 使用终端当前所在目录，而非脚本所在目录
STATS = Counter()
WARNINGS = []
HEADERS = ["No.", "Bk.D.", "Val.D.", "Type", "Vou.",
           "Vou.No./Trans.No./Details", "DebitAmount", "CreditAmount",
           "Balance", "ReferenceNo.", "Notes"]
RAW_FORMS = [
    ["Serial No. No.", "No."],
    ["Booked Date Bk.D.", "Bk.D."],
    ["Value Date Vou.", "Val.D."],
    ["Transaction Type Type", "Type"],
    ["Debit Vou.", "Vou."],
    ["Vou. No./Trans. No./Details Vou. No./Trans. No./Details", "Vou.No./Trans.No./Details"],
    ["Debit Amount Debit Amount", "DebitAmount"],
    ["Credit Amount Credit Amount", "CreditAmount"],
    ["Balance Balance", "Balance"],
    ["Institution/Cashier/Sequence Reference No.", "ReferenceNo.", "Reference No."],
    ["Notes Notes", "Notes"],
]

def norm(s):
    return re.sub(r'[^a-z0-9]', '', s.casefold())

HEADER_LOOKUP = {norm(s): i for i, forms in enumerate(RAW_FORMS) for s in forms}
ALIGNMENTS = [WD_ALIGN_PARAGRAPH.CENTER]*5 + [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.RIGHT]*3 + [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]

def has_image(p):
    return bool(p._p.xpath('.//w:drawing | .//w:pict | .//w:object'))

def safe_text_para(p):
    return not bool(p._p.xpath('.//w:drawing | .//w:pict | .//w:object | .//w:fldChar | .//w:instrText | .//w:hyperlink | .//w:bookmarkStart | .//w:bookmarkEnd | .//w:footnoteReference | .//w:endnoteReference'))

def child(parent, tag):
    obj = parent.find(qn(tag))
    if obj is None:
        obj = OxmlElement(tag)
        parent.append(obj)
    return obj

def fmt_run(run):
    # Any drawing-containing run is left byte-for-byte alone to protect image location.
    if run._r.xpath('.//w:drawing | .//w:pict | .//w:object'):
        STATS['image_runs_untouched'] += 1
        return
    rpr = run._element.get_or_add_rPr()
    fonts = child(rpr, 'w:rFonts')
    for attr, value in (('ascii', 'Times New Roman'), ('hAnsi','Times New Roman'),
                        ('eastAsia','SimSun'), ('cs','Times New Roman')):
        fonts.set(qn('w:'+attr),value)
    for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','csTheme'):
        fonts.attrib.pop(qn('w:'+attr), None)
    for tag,val in [('w:sz','16'),('w:szCs','16'),('w:w','100'),('w:spacing','0'),('w:position','0')]:
        child(rpr,tag).set(qn('w:val'), val)
    STATS['text_runs_formatted'] += 1

def fmt_para(p, alignment=None, indent_to_table=False):
    if has_image(p):
        # No alignment, indentation, spacing, breaks or run changes for image paragraphs.
        STATS['image_paragraphs_untouched'] += 1
        return
    for run in p.runs: fmt_run(run)
    pf = p.paragraph_format
    pf.space_before=Pt(0)
    pf.space_after=Pt(0)
    pf.line_spacing=Pt(8)
    pf.line_spacing_rule=WD_LINE_SPACING.EXACTLY
    if alignment is not None: p.alignment=alignment
    if indent_to_table:
        # Only when supplied a positive, verified main-table left/right margin.
        left,right=indent_to_table
        pf.left_indent=Pt(left)
        pf.right_indent=Pt(right)
    STATS['paragraphs_formatted'] += 1

def actual_cells(row):
    # enumerate actual w:tc elements, not repeated virtual cells for merged columns.
    for tc in row._tr.tc_lst:
        from docx.table import _Cell
        yield _Cell(tc, row.table)

def text_in_cell(cell):
    return ' '.join(p.text for p in cell.paragraphs).strip()

def header_hits(row):
    found={}
    for idx,cell in enumerate(actual_cells(row)):
        key=norm(text_in_cell(cell))
        if key in HEADER_LOOKUP:
            j=HEADER_LOOKUP[key]
            if j not in found:found[j]=(idx,cell)
    return found

def replace_header_cell(cell, s):
    if len(cell.paragraphs)!=1 or not safe_text_para(cell.paragraphs[0]):
        STATS['complex_headers_skipped'] += 1
        return
    p=cell.paragraphs[0]
    if p.text!=s:
        # Preserve paragraph properties, cell properties, other cell contents.
        runs=p.runs
        if runs:
            runs[0].text=s
            for r in runs[1:]:r.text=''
        else:p.add_run(s)
        STATS['headers_replaced']+=1

def structure_width_signature(tbl):
    # Assert that this code leaves all table grid and cell width elements untouched.
    from lxml import etree
    elems=tbl._tbl.xpath('.//w:tblGrid | .//w:gridCol | .//w:tcW | .//w:tblW | .//w:tblInd | .//w:gridSpan | .//w:gridBefore | .//w:gridAfter')
    return [etree.tostring(e) for e in elems]

def process_table(tbl, label):
    before=structure_width_signature(tbl)
    rows=list(tbl.rows)
    best={}; header_idx=None
    for ri,row in enumerate(rows[:min(9,len(rows))]):
        hits=header_hits(row)
        if len(hits)>len(best):best=hits;header_idx=ri
    # Require substantive header evidence; not other incidental tables.
    bank=(len(best)>=8 and all(i in best for i in (0,5,6,7,8,9,10)))
    if bank:
        STATS['bank_tables']+=1
        if len(best)!=11:WARNINGS.append(f'{label}: 仅识别 {len(best)}/11 项表头，其余未替换。')
        for j,(_,cell) in best.items():replace_header_cell(cell,HEADERS[j])
    # Only align rows with clearly 11 *physical* cells and a recognized header sequence.
    for ri,row in enumerate(rows):
        cells=list(actual_cells(row))
        aligned = bank and len(cells)==11 and (ri>=header_idx)
        for ci,cell in enumerate(cells):
            for p in cell.paragraphs:
                fmt_para(p, ALIGNMENTS[ci] if aligned else None)
            for nested_idx,nested in enumerate(cell.tables):
                process_table(nested,f'{label}/nested{nested_idx+1}')
    if before!=structure_width_signature(tbl):
        raise RuntimeError(f'{label}: 列宽/表格网格结构发生意外变化，已停止保存！')
    STATS['tables_scanned']+=1

def text_block_text(p):
    return norm(p.text)

def process_story(story, label, do_indents=False):
    # No deleting empty paragraphs, no removing page/section breaks, no changing images.
    for para in story.paragraphs:
        fmt_para(para)
    for ti,tbl in enumerate(story.tables,1):process_table(tbl,f'{label} Table {ti}')

def select_file():
    files=sorted((p for p in BASE.glob('*.docx') if not p.name.startswith('~$')
                  and not p.stem.endswith('_格式处理')),key=lambda x:x.name.casefold())
    if not files:raise FileNotFoundError(f'同目录无待处理 DOCX: {BASE}')
    print(f'当前目录：{BASE}')
    print('可处理 DOCX：')
    for i,p in enumerate(files,1):print(f'{i:2}  {p.name}')
    while True:
        s=input('输入文件编号（q退出）：').strip()
        if s.lower()=='q':return None
        if s.isdigit() and 1<=int(s)<=len(files):return files[int(s)-1]
        print('编号无效，请重试。')

def main():
    source=select_file()
    if not source:return
    output=source.with_name(source.stem+'_格式处理.docx')
    if output.exists():
        print('输出已存在，绝不覆盖：',output)
        return
    doc=Document(source)
    process_story(doc,'Body')
    # Unique header/footer parts only, avoid double-editing linked parts.
    visited=set()
    for sec in doc.sections:
        for name in ('header','footer','first_page_header','first_page_footer','even_page_header','even_page_footer'):
            story=getattr(sec,name)
            if story.is_linked_to_previous:continue
            identity=id(story._element)
            if identity not in visited:
                visited.add(identity)
                process_story(story,name)
    doc.save(output)
    log=output.with_suffix('.log.txt')
    log.write_text('只改文字，不改列宽、不碰图片段落、不删除交易数据。\n'+
                   '\n'.join(f'{k}: {v}' for k,v in sorted(STATS.items()))+
                   '\n警告：\n'+'\n'.join(WARNINGS or ['无'])+
                   '\n注意：普通文字排版可能改变分页。表格外缩进维持原状（无法仅凭文字推断主表左右边界）。\n',encoding='utf-8')
    print('\n完成：',output,'\n日志：',log)

if __name__=='__main__':main()